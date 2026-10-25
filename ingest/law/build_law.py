#!/usr/bin/env python3
"""Build data/law/*.jsonl, the Rule & law lens: forms of government, government spans per polity,
laws and policies, regime scores and Seshat law variables.

Usage: python3 ingest/law/build_law.py [RAW_DIR] [SESHAT_DIR]
RAW_DIR (default /mnt/project-files/atlas/raw/law) holds:
  p122.json                 Wikidata "basic form of government" (P122) with start/end qualifiers, for every
                            polity with a Wikidata id (SPARQL in README.md)
  iso.json                  polity Wikidata id -> ISO 3166 alpha-3 code (own P298, or P17 country's)
  iso_capitals.json         ISO3 -> [lat, lon, capital] from Wikidata (P36, P625)
  political-regime-lied.csv, universal-suffrage-lied.csv, universal-suffrage-women-lied.csv
                            Lexical Index of Electoral Democracy v6.9 via Our World in Data
SESHAT_DIR (default /mnt/project-files/atlas/raw/seshat) holds Equinox_data-1.0.zip and api/core_polities.jsonl.
"""
import csv, difflib, io, json, os, sys, zipfile
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
from forms import FAMILIES, FORMS, POLICIES, WIKIDATA, PRIORITY  # noqa: E402
from government import GOV  # noqa: E402
from laws import LAWS, SUFFRAGE_NOTES  # noqa: E402

RAW = (sys.argv[1] if len(sys.argv) > 1 else '/mnt/project-files/atlas/raw/law').rstrip('/') + '/'
SESHAT = (sys.argv[2] if len(sys.argv) > 2 else '/mnt/project-files/atlas/raw/seshat').rstrip('/') + '/'
OUT = os.path.join(ROOT, 'data', 'law')
WIKI = lambda a: {'cite': f'Wikipedia: {a}', 'url': 'https://en.wikipedia.org/wiki/' + a.replace(' ', '_')}
WD_SRC = {'cite': 'Wikidata, basic form of government (P122), retrieved 2026-10-24', 'url': 'https://www.wikidata.org/wiki/Property:P122'}
LIED_SRC = {'cite': 'Skaaning, Gerring and Bartusevicius, Lexical Index of Electoral Democracy (LIED) v6.9, processed by Our World in Data (2026)',
            'url': 'https://ourworldindata.org/grapher/political-regime-lexical'}
SESHAT_SRC = {'cite': 'Seshat Global History Databank, Equinox-2020 release (Turchin et al.), CC BY-NC-SA 4.0',
              'url': 'https://github.com/MajidBenam/Equinox_data'}
LIED_NAMES = ['non-electoral autocracy', 'one-party autocracy', 'multi-party autocracy without elected executive',
              'multi-party autocracy', 'exclusive democracy', 'male democracy', 'electoral democracy', 'full democracy (polyarchy)']
LAST = 2025


def slug_year(y):
    return f'bc{-y}' if y < 0 else str(y)


def rows(path):
    return [json.loads(l) for l in open(path, encoding='utf-8') if l.strip()]


polities = {}
for f in ('states-war/polities.jsonl', 'living-atlas/polities.jsonl', 'order-chaos/polities.jsonl'):
    p = os.path.join(ROOT, 'data', f)
    if os.path.exists(p):
        for r in rows(p):
            base = polities.setdefault(r['id'], {})
            for k, v in r.items():
                base.setdefault(k, v)
# A polity no other workstream has yet: classical Athens.
STUBS = [{'id': 'polity:athens', 'kind': 'polity', 'name': 'Athens', 'stub': True, 'start': -594, 'end': -322,
          'subtype': 'city_state', 'wikidata': 'Q844930', 'lat': 37.97, 'lon': 23.73, 'lenses': ['law', 'order'],
          'sources': [WIKI('Classical Athens')], 'status': 'draft'}]
for s in STUBS:
    polities.setdefault(s['id'], s)

form_family = {f[0]: f[2] for f in FORMS}
out = defaultdict(list)

# ---------- vocabulary ----------
for slug, name, color, what in FAMILIES:
    out['forms'].append({'id': f'idea:gov_family:{slug}', 'kind': 'idea', 'subtype': 'institution', 'name': name,
                         'what': what, 'color': color, 'lenses': ['law'], 'sources': [WIKI('List of forms of government')], 'status': 'draft'})
for slug, name, fam, what, example in FORMS:
    r = {'id': f'idea:gov:{slug}', 'kind': 'idea', 'subtype': 'institution', 'name': name, 'family': f'idea:gov_family:{fam}',
         'what': what, 'lenses': ['law'], 'sources': [WIKI('List of forms of government')], 'status': 'draft'}
    if example:
        r['example'] = example
    out['forms'].append(r)
for slug, name, what, why in POLICIES:
    out['forms'].append({'id': f'idea:policy:{slug}', 'kind': 'idea', 'subtype': 'institution', 'name': name, 'what': what,
                         'why': why, 'lenses': ['law'], 'sources': [WIKI(name if slug != 'limits_on_rulers' else 'Constitutionalism')], 'status': 'draft'})
out['forms'] += STUBS

# ---------- government spans: curated ----------
seen = set()
for pid, s, e, form, note, art in GOV:
    assert pid in polities, pid
    assert form in form_family, form
    rid = f'rel:gov:{pid.split(":", 1)[1]}:{slug_year(s)}'
    assert rid not in seen, rid
    seen.add(rid)
    out['government'].append({'id': rid, 'kind': 'relation', 'type': 'governed_as', 'from': pid, 'to': f'idea:gov:{form}',
                              'start': s, 'end': e, 'note': note, 'basis': 'curated', 'lenses': ['law'],
                              'sources': [WIKI(art)], 'status': 'draft'})

# ---------- government spans: Wikidata P122 ----------
qid_pol = defaultdict(list)
for pid, p in polities.items():
    if p.get('wikidata'):
        qid_pol[p['wikidata']].append(pid)


def wd_year(b, k):
    if k not in b:
        return None
    v = b[k]['value']
    neg = v.startswith('-')
    y = int(v.lstrip('-+').split('-')[0])
    return -y if neg else (y or None)


wd = defaultdict(dict)  # polity -> form -> (start, end, label)
for b in json.load(open(RAW + 'p122.json')):
    form = WIKIDATA.get(b['gLabel']['value'])
    if not form:
        continue
    q = b['p']['value'].rsplit('/', 1)[1]
    for pid in qid_pol.get(q, []):
        p = polities[pid]
        s, e = wd_year(b, 's'), wd_year(b, 'e')
        ps, pe = p.get('start'), p.get('end')
        # Qualifier dates belong to the Wikidata item, which may be wider than this Cliopatria slice: clip.
        s = max(s, ps) if s is not None and ps is not None else (s if s is not None else ps)
        e = min(e, pe) if e is not None and pe is not None else (e if e is not None else pe)
        if s is not None and e is not None and s > e:
            continue
        wd[pid].setdefault(form, (s, e, b['gLabel']['value'], 's' in b))
for pid in sorted(wd):
    forms = sorted(wd[pid], key=PRIORITY.index)
    for i, form in enumerate(forms):
        s, e, label, dated = wd[pid][form]
        r = {'id': f'rel:gov:wd:{pid.split(":", 1)[1]}:{form}', 'kind': 'relation', 'type': 'governed_as', 'from': pid,
             'to': f'idea:gov:{form}', 'basis': 'wikidata', 'rank': i, 'wikidata_label': label, 'lenses': ['law'],
             'sources': [{**WD_SRC, 'url': 'https://www.wikidata.org/wiki/' + polities[pid]['wikidata']}], 'status': 'sourced'}
        if s is not None:
            r['start'] = s
        if e is not None:
            r['end'] = e
        if not dated:
            r['note'] = 'Wikidata gives no dates for this form; it is applied to the whole life of the polity.'
        out['government'].append(r)

# ---------- regimes 1789+, Lexical Index ----------
# Keyed by OWID entity name, which covers historical states (Prussia, USSR, West Germany) as well as ISO countries.
code_ent = {}


def owid(name, col):
    d = defaultdict(dict)
    for r in csv.DictReader(open(RAW + name, encoding='utf-8')):
        if r[col] != '':
            d[r['entity']][int(r['year'])] = int(float(r[col]))
            if len(r['code']) == 3:
                code_ent[r['code']] = r['entity']
    return d


regime = owid('political-regime-lied.csv', 'regime_lied')
suff = owid('universal-suffrage-lied.csv', 'suffrage_lied')
fsuff = owid('universal-suffrage-women-lied.csv', 'female_suffrage_lied')
caps = json.load(open(RAW + 'iso_capitals.json'))
iso = json.load(open(RAW + 'iso.json'))
# Polities Wikidata can't place: the Lexical Index entity (or ISO3 code) that covers them.
MANUAL_ISO = {
    'polity:russian': 'RUS', 'polity:russian_republic': 'RUS', 'polity:ussr': 'USSR', 'polity:ottoman': 'Ottoman Empire',
    'polity:german_democratic_republic': 'East Germany', 'polity:kingdom_of_bavaria': 'Kingdom of Bavaria',
    'polity:kingdom_of_saxony': 'Kingdom of Saxony', 'polity:kingdom_of_wurttemberg': 'Kingdom of Wurttemberg',
    'polity:kingdom_of_hanover': 'Hanover', 'polity:grand_duchy_of_tuscany': 'Grand Duchy of Tuscany',
    'polity:democratic_republic_of_vietnam': 'Democratic Republic of Vietnam', 'polity:united_provinces_of_central_america': 'United Provinces of Central America',
    'polity:qing': 'CHN', 'polity:german_e': 'DEU', 'polity:weimar_republic': 'DEU', 'polity:nazi_germany': 'DEU',
    'polity:federal_republic_of_germany': 'West Germany', 'polity:north_german_confederation': 'Prussia', 'polity:kingdom_of_prussia': 'Prussia',
    'polity:austria_h': 'AUT', 'polity:austrian_empire': 'Austria-Hungary', 'polity:habsburg': 'Austria-Hungary', 'polity:second_republic_of_austria': 'AUT',
    'polity:british': 'GBR', 'polity:british_empire': 'GBR', 'polity:kingdom_of_great_britain': 'GBR',
    'polity:british_raj': 'IND', 'polity:mughal': 'IND', 'polity:empire_of_brazil': 'BRA', 'polity:mexico': 'MEX',
    'polity:centralist_mexico': 'MEX', 'polity:joseon': 'KOR', 'polity:korean_empire': 'KOR', 'polity:kingdom_of_portugal': 'PRT',
    'polity:kingdom_of_italy': 'ITA', 'polity:kingdom_of_sardinia': 'Kingdom of Sardinia', 'polity:dutch_republic': 'NLD',
    'polity:swedish_empire': 'SWE', 'polity:united_kingdoms_of_sweden_and_norway': 'SWE', 'polity:denmark_norway': 'DNK',
    'polity:denmark': 'DNK', 'polity:luxembourg': 'LUX', 'polity:hungary': 'HUN', 'polity:hungarian_people_s_republic': 'HUN',
    'polity:romania': 'ROU', 'polity:socialist_republic_of_romania': 'ROU', 'polity:kingdom_of_romania': 'ROU',
    'polity:syria': 'SYR', 'polity:second_syrian_republic': 'SYR', 'polity:eire': 'IRL', 'polity:irish_free_state': 'IRL',
    'polity:qajar_dynasty': 'IRN', 'polity:rattanakosin_kingdom': 'THA', 'polity:nguy_n_dynasty': 'VNM',
    'polity:kingdom_of_egypt': 'EGY', 'polity:kingdom_of_iraq': 'IRQ', 'polity:kingdom_of_afghanistan': 'AFG',
    'polity:republic_of_afghanistan': 'AFG', 'polity:mongolian_people_s_republic': 'MNG', 'polity:second_polish_republic': 'POL',
    'polity:kingdom_of_greece': 'GRC', 'polity:kingdom_of_bulgaria': 'BGR', 'polity:serbia': 'SRB', 'polity:kingdom_of_hejaz_and_nejd': 'SAU',
    'polity:emirate_of_nejd': 'SAU', 'polity:empire_of_haiti': 'HTI', 'polity:haiti': 'HTI', 'polity:gran_colombia': 'Great Colombia',
    'polity:new_granada': 'COL', 'polity:united_states_of_colombia': 'COL', 'polity:granadine_confederation': 'COL',
    'polity:zaire': 'COD', 'polity:rhodesia': 'ZWE', 'polity:zimbabwe_rhodesia': 'ZWE', 'polity:kingdom_of_libya': 'LBY',
    'polity:kingdom_of_laos': 'LAO', 'polity:people_s_republic_of_kampuchea': 'KHM', 'polity:united_arab_republic': 'EGY',
    'polity:kingdom_of_iceland': 'ISL', 'polity:kingdom_of_lithuania': 'LTU', 'polity:georgia': 'GEO', 'polity:ethiopia': 'ETH',
    'polity:french_first_republic': 'FRA', 'polity:french_directory': 'FRA', 'polity:french_consulate': 'FRA',
    'polity:first_french_empire': 'FRA', 'polity:kingdom_of_france': 'FRA', 'polity:bourbon_kingdom_of_france': 'FRA',
    'polity:french_second_republic': 'FRA', 'polity:vichy_france': 'FRA', 'polity:french_fifth_republic': 'FRA',
    'polity:beylik_of_tunis': 'TUN', 'polity:batavian_republic': 'NLD', 'polity:helvetic_republic': 'CHE',
    'polity:kingdom_of_the_two_sicilies': 'Kingdom of the Two Sicilies', 'polity:papal_states': 'ITA', 'polity:kingdom_of_spain': 'ESP',
    'polity:francoist_spain': 'ESP', 'polity:second_spanish_republic': 'ESP', 'polity:first_spanish_republic': 'ESP',
}
pol_iso = {}
for pid, p in polities.items():
    if (p.get('end') is not None and p['end'] < 1789) or p.get('start') is None:
        continue
    code = MANUAL_ISO.get(pid)
    if not code and p.get('wikidata') in iso:
        cands = iso[p['wikidata']]
        # Only the item's own country code: "country" (P17) also tags colonies with their metropole.
        own = [c for t, c in cands if t == 'own']
        code = own[0] if own else None
    code = code_ent.get(code, code)
    if code and code in regime:
        pol_iso[pid] = code
for pid, code in sorted(pol_iso.items()):
    p = polities[pid]
    s, e = max(p['start'], 1789), min(p['end'] if p.get('end') is not None else LAST, LAST)
    ys = [y for y in sorted(regime[code]) if s <= y <= e]
    if not ys:
        continue
    pts, last = [], None
    for y in ys:  # change points only; the value holds until the next point
        v = regime[code][y]
        if v != last:
            pts.append([y, v])
            last = v
    if pts[-1][0] != ys[-1]:
        pts.append([ys[-1], last])
    out['regimes'].append({'id': f'm:regime_lied:{pid.split(":", 1)[1]}:lied', 'kind': 'measure', 'of': pid, 'metric': 'regime_lied',
                           'unit': 'LIED class 0-7', 'step': True, 'country': code, 'classes': LIED_NAMES, 'points': pts,
                           'lenses': ['law'], 'sources': [LIED_SRC], 'status': 'sourced'})


def holder(code, y):
    """The polity that held this country in year y, preferring the shortest-lived (most specific) match."""
    c = [pid for pid, k in pol_iso.items() if k == code and polities[pid]['start'] <= y <= (polities[pid].get('end') or LAST)]
    return min(c, key=lambda pid: (polities[pid].get('end') or LAST) - polities[pid]['start']) if c else None


# Suffrage firsts: the year the index first switches on, if it was off the year before.
for policy, series, test in (('male_suffrage', suff, lambda v: v >= 1), ('female_suffrage', fsuff, lambda v: v >= 1)):
    for code, ent in sorted(code_ent.items()):
        if code not in caps:
            continue
        ser = series[ent]
        ys = sorted(ser)
        first = next((y for i, y in enumerate(ys) if i and test(ser[y]) and not test(ser[ys[i - 1]])), None)
        if not first:
            continue
        nm = ent
        who = 'women' if policy == 'female_suffrage' else 'all men'
        note, art = SUFFRAGE_NOTES.get((code, policy), (None, None))
        r = {'id': f'event:law:{policy}:{code.lower()}', 'kind': 'event', 'subtype': 'law',
             'name': f'{"Women" if policy == "female_suffrage" else "All men"} win the vote in {nm}', 'start': first,
             'lat': caps[code][0], 'lon': caps[code][1], 'policy': f'idea:policy:{policy}', 'effect': 'adopt', 'country': code,
             'what': note or f'From {first}, {who} who are citizens could vote in national elections in {nm} (Lexical Index coding).',
             'why': 'The index counts suffrage as universal when virtually all adult citizens of that sex may vote; informal barriers are not counted.',
             'lenses': ['law'], 'sources': [LIED_SRC] + ([WIKI(art)] if art else []), 'status': 'sourced'}
        h = holder(ent, first)
        if h:
            r['by'] = h
        out['laws'].append(r)

# ---------- curated laws ----------
for slug, name, yr, by, policy, effect, lat, lon, what, why, art in LAWS:
    approx = isinstance(yr, str)
    y = int(yr.lstrip('~')) if approx else yr
    assert by is None or by in polities, by
    r = {'id': f'event:law:{slug}', 'kind': 'event', 'subtype': 'law', 'name': name, 'start': y, 'lat': lat, 'lon': lon,
         'policy': f'idea:policy:{policy}', 'effect': effect, 'what': what, 'why': why, 'lenses': ['law'],
         'sources': [WIKI(art)], 'status': 'draft'}
    if approx:
        r['start_approx'] = True
    if by:
        r['by'] = by
    out['laws'].append(r)

# ---------- Seshat: courts, codes, divine rulers, limits on the executive ----------
SESHAT_VARS = {'Formal legal code': 'legal_code', 'Courts': 'courts', 'Judges': 'judges', 'Professional Lawyers': 'lawyers',
               'Full-time bureaucrats': 'bureaucrats', 'Rulers are gods': 'divine_rulers',
               'Rulers are legitimated by gods': 'divine_legitimation',
               'Constraint on executive by government': 'exec_constraint_gov',
               'Constraint on executive by non-government': 'exec_constraint_nongov'}
VAL = {'present': 1, 'inferred present': 1, 'absent': 0, 'inferred absent': 0}
zp = SESHAT + 'Equinox_data-1.0.zip'
if os.path.exists(zp):
    import openpyxl  # only this step needs it
    z = zipfile.ZipFile(zp)
    wb = openpyxl.load_workbook(io.BytesIO(z.read([n for n in z.namelist() if n.endswith('.xlsx')][0])), read_only=True)
    pr = list(wb['Polities'].iter_rows(values_only=True))
    eqp = {r[2]: (r[1], int(r[3]), int(r[4])) for r in pr[1:] if r[2]}
    core = rows(SESHAT + 'api/core_polities.jsonl')
    by_years = defaultdict(list)
    for c in core:
        try:
            by_years[(int(c['start_year']), int(c['end_year']))].append(c)
        except ValueError:
            pass
    sim = lambda a, b: difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()
    eq2core = {}
    for pol, (nm, s, e) in eqp.items():
        cands = by_years.get((s, e)) or core
        best = max(cands, key=lambda c: sim(nm, c['long_name']))
        if sim(nm, best['long_name']) >= (0.5 if cands is not core else 0.86):
            eq2core[pol] = best['name']
    s2p = defaultdict(list)
    for pid, p in polities.items():
        for s in (p.get('ext') or {}).get('seshat') or []:
            s2p[s].append(pid)
    facts = defaultdict(dict)  # (polity, metric) -> year -> value
    for r in wb['Equinox2020_CanonDat'].iter_rows(min_row=2, values_only=True):
        var, pol, val = r[4], r[1], r[5]
        if var not in SESHAT_VARS or val not in VAL or pol not in eq2core:
            continue
        _, s, e = eqp[pol]
        if r[7] is not None:
            try:
                s = int(r[7]); e = int(r[8]) if r[8] is not None else s
            except (TypeError, ValueError):
                pass
        for pid in s2p.get(eq2core[pol], []):
            d = facts[(pid, SESHAT_VARS[var])]
            d[s or 1] = VAL[val]
            d[e or 1] = VAL[val]
    for (pid, metric), d in sorted(facts.items()):
        pts = [[y, d[y]] for y in sorted(d)]
        out['seshat'].append({'id': f'm:{metric}:{pid.split(":", 1)[1]}:seshat', 'kind': 'measure', 'of': pid, 'metric': metric,
                              'unit': 'presence (1 present, 0 absent)', 'points': pts, 'lenses': ['law'],
                              'sources': [SESHAT_SRC], 'status': 'sourced'})

os.makedirs(OUT, exist_ok=True)
for name, rs in out.items():
    with open(os.path.join(OUT, name + '.jsonl'), 'w', encoding='utf-8') as f:
        for r in rs:
            f.write(json.dumps(r, ensure_ascii=False) + '\n')
    print(name, len(rs))
